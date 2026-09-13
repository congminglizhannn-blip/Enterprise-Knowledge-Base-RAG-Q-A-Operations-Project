from docx import Document as DocxDocument
from openpyxl import load_workbook


def parse_docx_text(path: str) -> str:
    doc = DocxDocument(path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def parse_xlsx_text(path: str) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        rows.append(f"工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)
