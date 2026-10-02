from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

from app.services.parser.office import parse_docx_text, parse_xlsx_text


def test_docx_parser_preserves_heading_and_table_order(tmp_path) -> None:
    document = DocxDocument()
    document.add_heading("付款流程", level=1)
    outlined = document.add_paragraph("自定义层级标题")
    outline_level = OxmlElement("w:outlineLvl")
    outline_level.set(qn("w:val"), "1")
    outlined._p.get_or_add_pPr().append(outline_level)
    numbered = document.add_paragraph("编号步骤")
    numbering_entry = document.part.numbering_part.element.findall(qn("w:num"))[0]
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), numbering_entry.get(qn("w:numId")))
    num_pr.extend([level, num_id])
    numbered._p.get_or_add_pPr().insert(0, num_pr)
    numbered_outline = OxmlElement("w:outlineLvl")
    numbered_outline.set(qn("w:val"), "2")
    numbered._p.get_or_add_pPr().append(numbered_outline)
    document.add_paragraph("先核验合同。")
    document.add_paragraph("嵌套条目", style="List Bullet 2")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "角色"
    table.cell(0, 1).text = "职责"
    table.cell(1, 0).text = "财务"
    table.cell(1, 1).text = "审核付款"
    path = tmp_path / "流程.docx"
    document.save(path)

    markdown = parse_docx_text(str(path))

    assert markdown.index("# 付款流程") < markdown.index("先核验合同。") < markdown.index("| 角色 | 职责 |")
    assert "## 自定义层级标题" in markdown
    assert any(marker + "编号步骤" in markdown for marker in ("1. ", "- "))
    assert "### 编号步骤" not in markdown
    assert "    - 嵌套条目" in markdown
    assert "| 财务 | 审核付款 |" in markdown


def test_xlsx_parser_renders_each_sheet_as_markdown_table(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "审批权限"
    sheet.append(["金额", "审批人"])
    sheet.append([10000, "财务经理"])
    path = tmp_path / "权限.xlsx"
    workbook.save(path)

    markdown = parse_xlsx_text(str(path))

    assert "## 审批权限" in markdown
    assert "| 金额 | 审批人 |" in markdown
    assert "| 10000 | 财务经理 |" in markdown


def test_docx_parser_extracts_embedded_excel_ole_workbook(tmp_path) -> None:
    document = DocxDocument()
    document.add_paragraph("角色职责")
    document.add_paragraph("点击图片可查看完整电子表格")
    path = tmp_path / "嵌入工作簿.docx"
    document.save(path)

    workbook = Workbook()
    workbook.active.append(["角色", "职责"])
    workbook.active.append(["项目经理", "组织立项论证"])
    embedded = BytesIO()
    workbook.save(embedded)
    malformed_dimensions = BytesIO()
    with ZipFile(embedded) as source, ZipFile(malformed_dimensions, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                assert 'ref="A1:B2"' in xml
                xml = xml.replace('ref="A1:B2"', 'ref="A1"', 1)
                data = xml.encode("utf-8")
            target.writestr(item, data)

    patched = tmp_path / "patched.docx"
    with ZipFile(path) as source, ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                xml = xml.replace("</w:body>", '<w:p><w:r><w:object><o:OLEObject r:id="rId50" ProgID="Excel.Sheet.12"/></w:object></w:r></w:p></w:body>', 1)
                data = xml.encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                xml = data.decode("utf-8").replace(
                    "</Relationships>",
                    '<Relationship Id="rId50" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="embeddings/book.xlsx"/></Relationships>',
                    1,
                )
                data = xml.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                xml = data.decode("utf-8").replace(
                    "</Types>",
                    '<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/></Types>',
                    1,
                )
                data = xml.encode("utf-8")
            target.writestr(item, data)
        target.writestr("word/embeddings/book.xlsx", malformed_dimensions.getvalue())

    markdown = parse_docx_text(str(patched))

    assert markdown.index("角色职责") < markdown.index("| 角色 | 职责 |")
    assert "| 项目经理 | 组织立项论证 |" in markdown
    assert "点击图片可查看完整电子表格" not in markdown
